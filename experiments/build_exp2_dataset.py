#!/usr/bin/env python3
"""연구데이터수집 시트(.xlsm) → 실험 2용 prompt-level CSV 변환.

입력:
  1_프롬프트      : item_id / domain / dataset_split / seed_prompt / v1_prompt / intent_summary / source_ref
  2_PII인스턴스   : item_id / pii_seq / pii_text / pii_type / final_tier / final_necessity / final_action

출력 (experiments/dataset.py::_load_csv_dataset 가 요구하는 스키마):
  id, split, domain, family_id, purpose, seed_prompt, v1_prompt, gold_entities_json
  + 참고용 test_category

오프셋(start/end)은 시트에 없다. v1_prompt 안에서 pii_text를 찾아 채운다.
같은 값이 여러 번 나오면 앞에서부터 아직 안 쓴 위치를 차례로 배정한다(occupancy).
찾지 못한 항목은 버리고 --report 에 사유와 함께 남긴다.

실행 (레포 루트에서):
    python experiments/build_exp2_dataset.py \
        --xlsm experiments/data/PersonaGuard_연구데이터수집_시트_2.xlsm \
        --outdir experiments/data \
        --report experiments/data/exp2_conversion_report.csv
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from pathlib import Path

import openpyxl

# ── 1. 매핑 표 ────────────────────────────────────────────────────────────────

# 시트 도메인(한글) → experiments/dataset.py DOMAINS
DOMAIN_MAP = {
    "건강·의료": "health",
    "법률·공공행정": "legal_admin",
    "채용·교육·커리어": "career_edu",
    "여행·지역추천": "travel_local",
    "금융·소비자·고객지원": "finance_cs",
}

# 시트 pii_type(한글) → PersonaGuard KNOWN_ENTITY_TYPES.
# None = PersonaGuard develop 이 탐지하지 못하는 타입 → gold 에서 제외한다
# (dataset.py 모듈 docstring: "탐지하지 못하는 항목은 gold_entities에 넣지 않는다").
# run_dev_v1_evaluation.py 의 GOLD_TYPE_MAP 과 같은 판단을 따른다.
TYPE_MAP = {
    "이름": "PERSON",
    "닉네임": "PERSON",          # GOLD_TYPE_MAP: PS_NICKNAME → PERSON
    "주소": "ADDRESS",
    "회사명_직장": "ORGANIZATION",
    "학교명": "ORGANIZATION",
    "전화번호": "PHONE",
    "이메일": "EMAIL",
    "계좌번호": "BANK_ACCOUNT",
    "카드번호": "CARD",
    "주민등록번호": "RRN",
    "여권번호": "PASSPORT",
    "운전면허번호": "DRIVER_LICENSE",
    # ── PersonaGuard 미지원 (GOLD_TYPE_MAP 에서 None 인 것들과 동일)
    "계정ID": None,              # PS_ID
    "생년월일": None,            # DT_BIRTH
    "나이": None,                # QT_AGE
    "혈액형": None,              # TM_BLOOD_TYPE
    "차량번호": None,            # QT_PLATE_NUMBER
    "전공": None,                # FD_MAJOR
    "직급": None,                # CV_POSITION
    "사건번호": None,            # QT_CASE_NUMBER
}

# challenge 시트의 변형 표기 → 기저 타입. 괄호 안이 은닉/분산 방식이다.
CHALLENGE_SUFFIXES = ("(공백 은닉)", "(분산", "(구 단위만)")

NECESSITY_MAP = {"필수": "essential", "조건부": "optional", "부수적": "unnecessary"}
ACTION_MAP = {"삭제": "delete", "마스킹": "mask", "일반화": "generalize", "보존": "keep"}
# GoldEntity.risk 는 1(낮음)~5(높음) 정수. 시트는 3단계 tier 이므로 5/3/1 로 편다.
TIER_TO_RISK = {"높음": 5, "중간": 3, "낮음": 1}

# "분산"(여러 문장 결합) 변형 표기 전용 처리.
# 시트의 pii_text는 조각을 합친 값(예: "010-0000-5008")이라 locate()가 v1_prompt
# 안에서 문자열 그대로 찾지 못해 전부 드롭되는 문제가 있었다(2026-08-19 발견,
# challenge 20건). v1_prompt 문장 자체는 고정 템플릿으로 조각을 명시적으로
# 알려주므로("앞부분은 X입니다. 뒷부분은 Y예요." 등), 그 문장에서 조각을 직접
# 정규식으로 뽑아 각 조각을 별도 gold entity로 등록한다. risk/necessity/action은
# 원래 그 PII 인스턴스 행의 final_* 라벨을 그대로 모든 조각에 적용한다(사람이
# "이 정보"를 하나의 단위로 판단한 것이므로 조각마다 새로 판단하지 않음).
FRAGMENT_PATTERNS: dict[str, "re.Pattern"] = {}


def _register_fragment_patterns():
    import re as _re
    FRAGMENT_PATTERNS["EMAIL"] = _re.compile(
        r"이메일 아이디는\s*(\S+?)입니다\. 도메인은\s*(\S+?)예요")
    FRAGMENT_PATTERNS["RRN"] = _re.compile(
        r"주민번호 앞자리는\s*(\d+)입니다\. 뒷자리는\s*(\d+)예요")
    FRAGMENT_PATTERNS["PHONE"] = _re.compile(
        r"연락처 앞부분은\s*([\d\-]+)입니다\. 뒷부분은\s*(\d+)예요")
    FRAGMENT_PATTERNS["PASSPORT"] = _re.compile(
        r"여권번호의 앞 네 자리는\s*(\S+?)입니다\. 가운데 문자는\s*(\S+?)예요\. "
        r"마지막 네 자리는\s*(\S+?)입니다")


_register_fragment_patterns()


def locate_split_fragments(
    v1: str, etype: str, taken: list[tuple[int, int]]
) -> list[tuple[str, int, int]] | None:
    """'분산' 변형의 v1_prompt 문장에서 조각들을 직접 찾는다.

    실패(패턴 미매칭 또는 조각 위치를 못 찾음)하면 None을 돌려주고, 호출부는
    기존처럼 drop() 처리한다 — 이 함수는 알려진 4개 템플릿에만 적용되고, 새로운
    템플릿이 추가되면 여기도 같이 갱신해야 한다.
    """
    pattern = FRAGMENT_PATTERNS.get(etype)
    if pattern is None:
        return None
    m = pattern.search(v1)
    if not m:
        return None
    result = []
    local_taken = list(taken)
    for frag in m.groups():
        span = locate(v1, frag, local_taken)
        if span is None:
            return None
        local_taken.append(span)
        result.append((frag, span[0], span[1]))
    return result


# 시트 challenge_type → dataset.py CHALLENGE_TRAPS
TRAP_MAP = {
    "공백_특수문자_은닉": "obfuscation",
    "여러문장_분산": "split_across_sentences",
    "간접정보_조합": "aggregation",
    "표_코드_로그_내부": "embedded_in_code",
    "인젝션_다중PII": "adversarial_instruction",
}

# 시트 pii_condition → dataset.py TEST_CATEGORIES
CONDITION_MAP = {"없음": "no_pii", "부수적": "optional_pii", "필수": "essential_pii", "혼합": "mixed_pii"}


def base_type(raw: str) -> tuple[str, str | None]:
    """'이메일(공백 은닉)' → ('이메일', '공백 은닉'). 변형 표기가 없으면 (raw, None)."""
    if "(" in raw and raw.endswith(")"):
        head, _, tail = raw.partition("(")
        return head.strip(), tail[:-1].strip()
    return raw.strip(), None


# ── 2. 시트 읽기 ──────────────────────────────────────────────────────────────


def read_sheets(xlsm: Path):
    wb = openpyxl.load_workbook(xlsm, data_only=True, read_only=True)

    ws = wb["1_프롬프트"]
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(hdr) if h}
    prompts = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        prompts.append({k: row[i] for k, i in idx.items()})

    ws2 = wb["2_PII인스턴스"]
    hdr2 = [c.value for c in next(ws2.iter_rows(min_row=1, max_row=1))]
    idx2 = {h: i for i, h in enumerate(hdr2) if h}
    inst = defaultdict(list)
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        rec = {k: row[i] for k, i in idx2.items()}
        iid = str(rec["item_id"]).strip()
        if "\n" in iid or "=" in iid:      # 시트 하단의 집계 수식 줄
            continue
        inst[iid].append(rec)
    wb.close()
    return prompts, inst


# ── 3. 오프셋 부여 ────────────────────────────────────────────────────────────


def locate(text: str, needle: str, taken: list[tuple[int, int]]) -> tuple[int, int] | None:
    """text 안에서 needle 을, 이미 배정된 구간(taken)과 겹치지 않는 첫 위치에 배정한다."""
    pos = 0
    while True:
        i = text.find(needle, pos)
        if i < 0:
            return None
        j = i + len(needle)
        if not any(s < j and i < e for s, e in taken):
            return i, j
        pos = i + 1


# ── 4. 변환 ──────────────────────────────────────────────────────────────────


def convert(prompts, inst):
    cases, dropped = [], []
    for p in prompts:
        iid = str(p["item_id"]).strip()
        split = (p.get("dataset_split") or "").strip()
        v1 = p.get("v1_prompt") or ""
        seed = p.get("seed_prompt") or ""
        domain_ko = (p.get("domain") or "").strip()
        domain = DOMAIN_MAP.get(domain_ko)
        if domain is None:
            dropped.append({"item_id": iid, "pii_seq": "", "pii_text": "", "pii_type": "",
                            "reason": f"미매핑 도메인 {domain_ko!r} — 문항 전체 제외"})
            continue

        golds, taken = [], []
        for rec in sorted(inst.get(iid, []), key=lambda r: int(r["pii_seq"] or 0)):
            raw_type = str(rec.get("pii_type") or "").strip()
            text = str(rec.get("pii_text") or "").strip()
            seq = rec.get("pii_seq")
            bt, variant = base_type(raw_type)

            def drop(reason):
                dropped.append({"item_id": iid, "pii_seq": seq, "pii_text": text,
                                "pii_type": raw_type, "reason": reason})

            if not text or bt == "(PII 없음)":
                continue
            if bt not in TYPE_MAP:
                drop(f"매핑표에 없는 타입 {bt!r}")
                continue
            etype = TYPE_MAP[bt]
            if etype is None:
                drop(f"{bt} — PersonaGuard 미지원 타입 (GOLD_TYPE_MAP=None)")
                continue

            nec = NECESSITY_MAP.get(str(rec.get("final_necessity") or "").strip())
            act = ACTION_MAP.get(str(rec.get("final_action") or "").strip())
            risk = TIER_TO_RISK.get(str(rec.get("final_tier") or "").strip())
            if nec is None or act is None or risk is None:
                drop(f"final 라벨 미기입/미매핑 (tier={rec.get('final_tier')!r} "
                     f"necessity={rec.get('final_necessity')!r} action={rec.get('final_action')!r})")
                continue

            span = locate(v1, text, taken)
            if span is None:
                fragments = locate_split_fragments(v1, etype, taken) if variant else None
                if fragments is None:
                    drop(f"v1_prompt 안에서 값을 찾지 못함"
                         + (f" (변형 표기: {variant})" if variant else ""))
                    continue
                for frag_text, fstart, fend in fragments:
                    taken.append((fstart, fend))
                    golds.append({
                        "text": frag_text, "type": etype, "start": fstart, "end": fend,
                        "risk": risk, "necessity": nec, "action": act,
                    })
                continue
            taken.append(span)
            golds.append({
                "text": text, "type": etype, "start": span[0], "end": span[1],
                "risk": risk, "necessity": nec, "action": act,
            })

        golds.sort(key=lambda g: g["start"])
        cases.append({
            "id": iid,
            "split": split,
            "domain": domain,
            "family_id": str(p.get("source_ref") or iid),
            "purpose": p.get("intent_summary") or "",
            "seed_prompt": seed,
            "v1_prompt": v1,
            "gold_entities_json": json.dumps(golds, ensure_ascii=False),
            "test_category": CONDITION_MAP.get(str(p.get("pii_condition") or "").strip(), ""),
            "challenge_trap": TRAP_MAP.get(str(p.get("challenge_type(Challenge만)") or "").strip(), ""),
        })
    return cases, dropped


FIELDS = ["id", "split", "domain", "family_id", "purpose", "seed_prompt",
          "v1_prompt", "gold_entities_json", "test_category", "challenge_trap"]


def write_csv(path: Path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsm", required=True, type=Path)
    ap.add_argument("--outdir", required=True, type=Path)
    ap.add_argument("--report", type=Path, default=None)
    args = ap.parse_args()

    prompts, inst = read_sheets(args.xlsm)
    cases, dropped = convert(prompts, inst)

    by_split = defaultdict(list)
    for c in cases:
        by_split[c["split"]].append(c)

    for split, rows in sorted(by_split.items()):
        out = args.outdir / f"PersonaGuard_exp2_{split}.csv"
        write_csv(out, rows)
        n_ent = sum(len(json.loads(r["gold_entities_json"])) for r in rows)
        print(f"[완료] {out}  {len(rows)}문항 / gold {n_ent}개")

    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        with args.report.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["item_id", "pii_seq", "pii_text", "pii_type", "reason"])
            w.writeheader()
            w.writerows(dropped)
        print(f"[제외 {len(dropped)}건] {args.report}")


if __name__ == "__main__":
    main()
